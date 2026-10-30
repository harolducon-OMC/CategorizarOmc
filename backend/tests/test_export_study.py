import io
import os
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core.export import build_study_categorization_excel


class StudyCategorizationExportTests(unittest.TestCase):
    def test_export_has_three_clean_sheets_and_keeps_blank_answers_blank(self):
        source = pd.DataFrame({
            "respondent_id": [1, 2],
            "CE1": ["Marca A", "Marca B"],
            "CE1_1": ["me gusta", None],
        })
        classifications = [
            {"question": "CE1", "kind": "brands", "labels": [["Marca A"], ["Marca B"]],
             "categories": ["Marca A", "Marca B"], "codeframe": [],
             "config": {"known_brands": ["Marca A", "Marca B"]},
             "brand_mapping": {"Marca A": "Marca A", "Marca B": "Marca B"}},
            {"question": "CE1_1", "kind": "sentiment", "labels": [["Positivo"], ["NS/NR"]],
             "categories": ["Positivo", "Neutral", "Negativo", "Mixto/ambivalente"],
             "codeframe": [{"nombre": "Positivo", "descripcion": "Favorable"}]},
        ]
        content = build_study_categorization_excel(source, classifications)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        self.assertEqual(workbook.sheetnames, ["Datos", "Preguntas y categorización", "Diccionario"])

        data = workbook["Datos"]
        self.assertEqual(next(data.iter_rows(values_only=True)), ("respondent_id", "CE1", "CE1_1"))
        self.assertEqual(next(data.iter_rows(min_row=2, max_row=2, values_only=True)), (1, "Marca A", "me gusta"))
        blank_source_row = next(data.iter_rows(min_row=3, max_row=3, values_only=True))
        self.assertEqual(tuple(blank_source_row) + (None,) * (3 - len(blank_source_row)), (2, "Marca B", None))

        coded = workbook["Preguntas y categorización"]
        self.assertEqual(next(coded.iter_rows(values_only=True)), (
            "Fila de encuesta", "CE1 | Respuesta", "CE1 | Categoría(s)",
            "CE1_1 | Respuesta", "CE1_1 | Categoría(s)",
        ))
        self.assertEqual(next(coded.iter_rows(min_row=2, max_row=2, values_only=True)),
                         (2, "Marca A", "Marca A", "me gusta", "Positivo"))
        # Aunque la clasificación antigua almacenase NS/NR, la celda fuente vacía se exporta vacía.
        blank_category_row = next(coded.iter_rows(min_row=3, max_row=3, values_only=True))
        self.assertEqual(tuple(blank_category_row) + (None,) * (5 - len(blank_category_row)),
                         (3, "Marca B", "Marca B", None, None))

        dictionary = workbook["Diccionario"]
        dictionary_rows = list(dictionary.iter_rows(values_only=True))
        self.assertEqual(dictionary_rows[0], (
            "Pregunta", "Tipo de análisis", "Código / categoría", "Descripción",
            "Palabras clave / sinónimos", "Alias / variantes",
        ))
        self.assertIn(("CE1", "Marcas", "Marca A", "Marca del catálogo permitido."),
                      [tuple(row[:4]) for row in dictionary_rows])
        workbook.close()


if __name__ == "__main__":
    unittest.main()
