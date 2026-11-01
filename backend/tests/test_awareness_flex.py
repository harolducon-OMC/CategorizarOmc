import io
import unittest

import pandas as pd
from openpyxl import load_workbook

from core import awareness_flex


class AwarenessFlexTests(unittest.TestCase):
    def setUp(self):
        self.frame = pd.DataFrame({
            "F1 Sexo": ["F", "M", "F"],
            "P1 Marca": ["Alfa", "Beta", "ninguna"],
            "P1A.1 Otras": ["Beta", "", "Alfa"],
            "P2.1 Alfa": ["1", "", "1"],
        })
        self.config = awareness_flex.parse_config({
            "entities": ["Alfa", "Beta"],
            "aliases": {"Alfa": ["alfa"], "Beta": ["beta"]},
            "demographics": {"sexo": "F1"},
            "analyses": [{"name": "Test", "tom": "P1 ", "esp": "P1A", "ayu": "P2"}, {"name": "Off"}],
        })

    def test_presets_have_three_modes(self):
        self.assertEqual(set(awareness_flex.presets()), {"bancos", "conglomerados", "personalizado"})

    def test_validation_detects_columns(self):
        rows = awareness_flex.detect_columns(list(self.frame.columns), self.config)
        tom = next(r for r in rows if r["campo"] == "TOM")
        self.assertEqual(tom["columnas"], "P1 Marca")

    def test_report_is_valid_xlsx_with_expected_sheets(self):
        data = awareness_flex.generate_report(self.frame, self.config)
        workbook = load_workbook(io.BytesIO(data))
        self.assertIn("Resumen Demográficos", workbook.sheetnames)
        self.assertIn("Test", workbook.sheetnames)
        self.assertNotIn("Off", workbook.sheetnames)

    def test_requires_entities_and_one_active_analysis(self):
        with self.assertRaises(ValueError):
            awareness_flex.parse_config({"entities": []})
        config = awareness_flex.parse_config({"entities": ["A"]})
        with self.assertRaises(ValueError):
            awareness_flex.generate_report(self.frame, config)


class FuzzyAndCoverageTests(unittest.TestCase):
    def setUp(self):
        answers = (["Adidas"] * 6 + ["Nike"] * 3 + ["Decatlon"] * 2 + ["Decathlon"] + ["Sportfitness"] * 2
                   + ["Puma, Nike"] + ["Ninguna", "", "Forma Boots"])
        self.frame = pd.DataFrame({"P1 Marca": answers})
        self.base = {
            "entities": ["Decathlon", "Sport fitness", "Forma", "Ninguna"],
            "aliases": {"Decathlon": ["decathlon"], "Sport fitness": ["sport fitness"], "Forma": ["forma"]},
            "analyses": [{"name": "A", "tom": "P1", "esp": "0", "ayu": "0"}],
        }

    def test_typos_are_grouped_with_the_brand(self):
        config = awareness_flex.parse_config(self.base)
        rows = {r["Entidad"]: r for r in awareness_flex.preview_results(self.frame, config)["analyses"][0]["rows"]}
        self.assertEqual(rows["Decathlon"]["TOM"], 3)        # Decathlon + 2 x "Decatlon"
        self.assertEqual(rows["Sport fitness"]["TOM"], 2)    # "Sportfitness"

    def test_fuzzy_can_be_turned_off(self):
        config = awareness_flex.parse_config({**self.base, "fuzzy": False})
        rows = {r["Entidad"]: r for r in awareness_flex.preview_results(self.frame, config)["analyses"][0]["rows"]}
        self.assertEqual(rows["Decathlon"]["TOM"], 1)

    def test_reserved_labels_are_not_entities(self):
        config = awareness_flex.parse_config(self.base)
        self.assertNotIn("Ninguna", config["entities"])
        self.assertEqual(config["ignored_entities"], ["Ninguna"])

    def test_coverage_reports_unrecognized_brands(self):
        config = awareness_flex.parse_config(self.base)
        coverage = awareness_flex.preview_results(self.frame, config)["analyses"][0]["coverage"]
        top = {item["respuesta"]: item["n"] for item in coverage["unrecognized"]}
        self.assertEqual(top["Adidas"], 6)
        self.assertEqual(top["Nike"], 3)
        self.assertLess(coverage["recognized"], coverage["answers"])

    def test_suggested_entities_group_variants(self):
        config = awareness_flex.parse_config(self.base)
        suggestion = awareness_flex.suggest_entities(self.frame, config)
        self.assertEqual(suggestion["entities"][0], "Adidas")
        self.assertEqual(sum(1 for name in suggestion["entities"] if name.lower().startswith("decat")), 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
