# -*- coding: utf-8 -*-
"""
core/awareness.py
Lógica de Awareness / Recordación (TOM, Espontáneo, Ayudado).

Extraído del repo https://github.com/HaroldDuconAbril/awareness
(sin interfaz Streamlit). Cambios respecto al original:
  - Se eliminó todo el código de interfaz (vive en pages/3_Awareness.py).
  - Se corrigió el .title() que convertía "BBVA" en "Bbva" en la tabla TOM.
"""

import io
import re
import unicodedata
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================
# CONFIGURACIONES PREDEFINIDAS
# ============================================================

# --- PRESET 1: BANCOS ---
BANKS = [
    "Banco Agrario", "Bancolombia", "Davivienda", "Banco de Bogotá",
    "BBVA", "Banco Caja Social", "Banco Popular", "Bancamía",
    "Banco de Occidente", "Banco Mundo Mujer",
]

BANK_ALIASES = {
    "Banco Agrario": ["banco agrario", "banco agrario de colombia", "agrario"],
    "Bancolombia": ["bancolombia"],
    "Davivienda": ["davivienda"],
    "Banco de Bogotá": ["banco de bogota", "banco de bogotá", "banco bogota", "banco bogotá"],
    "BBVA": ["bbva", "bbvva", "bvva", "bbwa", "bva", "bvvwa"],
    "Banco Caja Social": ["banco caja social", "caja social", "cajas social", "caja sosial"],
    "Banco Popular": ["banco popular", "popular"],
    "Bancamía": ["bancamia", "bancamía", "banca mia", "banca mía"],
    "Banco de Occidente": ["banco de occidente", "occidente"],
    "Banco Mundo Mujer": ["banco mundo mujer", "mundo mujer", "banco de la mujer", "fundacion de la mujer", "de la mujer"],
}

BANK_NORMALIZATIONS = [
    {"pattern": r"(?i).*Banco\s+de\s+Bogota.*", "replacement": "Banco de Bogotá"},
    {"pattern": r"(?i).*(?<!Banco\s)Caja\s+Social.*", "replacement": "Banco Caja Social"},
    {"pattern": r"(?i).*Bancamia.*", "replacement": "Bancamía"},
]

# --- PRESET 2: CONGLOMERADOS FINANCIEROS Y MARCAS GENERALES ---
CONGLOMERATES = [
    "Conglomerado BBVA", "Grupo cooperativo Coomeva", "Fundación Grupo Social",
    "Grupo Bolivar", "Conglomerado financiero Sura-Bancolombia", "Grupo Aval",
    "GNB Sudameris", "Conglomerado Credicorp capital"
]

CONGLOMERATE_ALIASES = {
    "Conglomerado BBVA": ["conglomerado bbva", "bbva"],
    "Grupo cooperativo Coomeva": ["grupo cooperativo coomeva", "coomeva"],
    "Fundación Grupo Social": ["fundacion grupo social", "grupo social"],
    "Grupo Bolivar": ["grupo bolivar", "bolivar"],
    "Conglomerado financiero Sura-Bancolombia": ["conglomerado financiero sura bancolombia", "sura", "bancolombia"],
    "Grupo Aval": ["grupo aval", "aval"],
    "GNB Sudameris": ["gnb sudameris", "sudameris", "gnb"],
    "Conglomerado Credicorp capital": ["conglomerado credicorp capital", "credicorp"]
}

CONGLOMERATE_NORMALIZATIONS = [
    {"pattern": r"(?i).*grupo\s+bancolombia.*", "replacement": "Grupo Bancolombia"},
    {"pattern": r"(?i).*(?<!grupo\s)bancolombia.*", "replacement": "Bancolombia"},
    {"pattern": r"(?i).*(exito|éxito).*", "replacement": "Grupo exito"},
    {"pattern": r"(?i).*ecopetrol.*", "replacement": "Grupo Ecopetrol"},
    {"pattern": r"(?i).*nutresa.*", "replacement": "Grupo nutresa"},
    {"pattern": r"(?i).*(sudameris|gnb).*", "replacement": "GNB Sudameris"},
    {"pattern": r"(?i).*credicorp.*", "replacement": "Conglomerado Credicorp capital"},
    {"pattern": r"(?i).*skandia.*", "replacement": "Conglomerado Skandia"},
    {"pattern": r"(?i).*(adidas|afidas|asidas|adida).*", "replacement": "Adidas"},
    {"pattern": r"(?i).*(coca-cola|cocacola|coca\s+cola).*", "replacement": "Coca-Cola"},
    {"pattern": r"(?i).*(av\s+villas|avvillas).*", "replacement": "Banco AV Villas"},
    {"pattern": r"(?i).*(nestle|nestlé).*", "replacement": "Nestlé"},
    {"pattern": r"(?i).*compensar.*", "replacement": "Compensar"}
]

NEGATIVOS = {
    "", "0", "no", "nan", "none", "false", "ninguno", "ninguna",
    "no se", "no sé", "no recuerdo", "ningun otro", "ningún otro", "no aplica",
}


# ============================================================
# UTILIDADES DE TEXTO
# ============================================================

def norm(value) -> str:
    if pd.isna(value): return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-z0-9\s]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()

def clean_col(value) -> str:
    if pd.isna(value): return ""
    text = str(value).strip().lower()
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text).strip()

def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", " ", str(name)).strip()
    return cleaned[:31] or "Hoja"

def apply_normalizations(value, normalizations: List[Dict]) -> str:
    if pd.isna(value):
        return value
    text = str(value).strip()
    for item in normalizations or []:
        pattern = item.get("pattern", "")
        replacement = item.get("replacement", "")
        if not pattern: continue
        try:
            if re.search(pattern, text):
                return replacement
        except re.error:
            pass
    return text

def _edit_distance_at_most_one(left: str, right: str) -> bool:
    """Return whether two strings differ by at most one insertion/deletion/substitution."""
    if left == right:
        return True
    if abs(len(left) - len(right)) > 1:
        return False

    left_index = right_index = differences = 0
    while left_index < len(left) and right_index < len(right):
        if left[left_index] == right[right_index]:
            left_index += 1
            right_index += 1
            continue
        differences += 1
        if differences > 1:
            return False
        if len(left) > len(right):
            left_index += 1
        elif len(right) > len(left):
            right_index += 1
        else:
            left_index += 1
            right_index += 1

    if left_index < len(left) or right_index < len(right):
        differences += 1
    return differences <= 1

def _contains_compact_alias(text: str, alias: str) -> bool:
    """Match spacing variants and one-character typos in a bounded word phrase."""
    alias_compact = alias.replace(" ", "")
    if not alias_compact:
        return False

    words = text.split()
    for start in range(len(words)):
        candidate = ""
        for end in range(start, min(len(words), start + 5)):
            candidate += words[end]
            if candidate == alias_compact:
                return True
            # Avoid fuzzy matches for short/common names where false positives are likelier.
            if (
                len(alias_compact) >= 8
                and abs(len(candidate) - len(alias_compact)) <= 1
                and _edit_distance_at_most_one(candidate, alias_compact)
            ):
                return True
    return False

def contains_entity(value, entity: str, aliases: Dict[str, List[str]]) -> bool:
    text = norm(value)
    if text in NEGATIVOS: return False
    for alias in aliases.get(entity, [entity]):
        alias_norm = norm(alias)
        if alias_norm and re.search(r"(^|\s)" + re.escape(alias_norm) + r"(\s|$)", text):
            return True
        if alias_norm and _contains_compact_alias(text, alias_norm):
            return True
    return False


# ============================================================
# DETECCIÓN DE COLUMNAS
# ============================================================

def find_col(df: pd.DataFrame, prefix: str) -> Optional[str]:
    prefix_norm = clean_col(prefix)
    if not prefix_norm or prefix_norm == "0": return None
    for column in df.columns:
        if clean_col(column) == prefix_norm:
            return column
    for column in df.columns:
        if _is_question_column(column, prefix_norm):
            return column
    return None

def _is_question_column(column, prefix_norm: str) -> bool:
    """Match a question code and its delimited subcolumns, never a longer code."""
    col_norm = clean_col(column)
    if col_norm == prefix_norm:
        return True
    return (
        col_norm.startswith(prefix_norm)
        and len(col_norm) > len(prefix_norm)
        and col_norm[len(prefix_norm)] in "_.- "
    )

def _question_column_sort_key(column, prefix_norm: str):
    col_norm = clean_col(column)
    exact = col_norm == prefix_norm
    suffix = col_norm[len(prefix_norm):]
    natural_suffix = tuple(
        (0, int(part)) if part.isdigit() else (1, part)
        for part in re.split(r"(\d+)", suffix)
    )
    return (0 if exact else 1, natural_suffix)

def prefixes_to_list(text: str) -> List[str]:
    lines = str(text).replace(",", "\n").splitlines()
    return [line.strip() for line in lines if line.strip()]

def find_cols(df: pd.DataFrame, prefixes_text: str) -> List[str]:
    detected = []
    raw_prefixes = [p for p in prefixes_to_list(prefixes_text) if p.strip() != "0"]
    prefixes = [clean_col(p) for p in raw_prefixes]
    if not prefixes: return detected
    for prefix in prefixes:
        has_exact_column = any(clean_col(column) == prefix for column in df.columns)
        matched = []
        for column in df.columns:
            col_norm = clean_col(column)
            is_exact = col_norm == prefix
            has_subquestion_suffix = (
                col_norm.startswith(prefix)
                and len(col_norm) > len(prefix)
                and col_norm[len(prefix)] in "_.-"
            )
            # When the exact code exists, a space usually starts a description,
            # not a split column. Without an exact code, allow headers like "P1 Marca".
            is_named_header = not has_exact_column and _is_question_column(column, prefix)
            if is_exact or has_subquestion_suffix or is_named_header:
                matched.append(column)
        for column in sorted(matched, key=lambda value: _question_column_sort_key(value, prefix)):
            if column not in detected:
                detected.append(column)
    return detected

def _is_empty_response(value) -> bool:
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    if isinstance(value, str):
        return value.strip() in ("", "0")
    try:
        return bool(value == 0)
    except (TypeError, ValueError):
        return False

def first_non_empty_response(df: pd.DataFrame, columns: List[str]) -> pd.Series:
    """Return one response per respondent, taking the first answered split column."""
    result = pd.Series([None] * len(df), index=df.index, dtype=object)
    filled = pd.Series(False, index=df.index)
    for column in columns:
        values = df[column]
        take = (~filled) & (~values.map(_is_empty_response))
        result.loc[take] = values.loc[take]
        filled = filled | take
    return result

def frame_with_categorized_responses(df: pd.DataFrame, classifications) -> pd.DataFrame:
    """Replace classified source columns with their saved, human-corrected labels."""
    output = df.copy()
    for item in classifications:
        if isinstance(item, dict):
            question = item.get("question")
            labels = item.get("labels") or []
        else:
            question = getattr(item, "question", None)
            labels = getattr(item, "labels", None) or []
        if question not in output.columns:
            continue
        raw_responses = df[question].tolist()
        cleaned = []
        for index, response in enumerate(raw_responses):
            if _is_empty_response(response) or index >= len(labels):
                cleaned.append(None)
                continue
            row_labels = labels[index] or []
            row_labels = [row_labels] if isinstance(row_labels, str) else row_labels
            values = [str(value).strip() for value in row_labels if value is not None and str(value).strip()]
            cleaned.append("; ".join(dict.fromkeys(values)) or None)
        output[question] = cleaned
    return output

def entity_from_aided_col(column_name: str, entities: List[str]) -> Optional[str]:
    text = norm(str(column_name).replace("\n", " "))
    for entity in entities:
        if norm(entity) in text: return entity
    return None


# ============================================================
# CÁLCULO DE AWARENESS Y TABLAS
# ============================================================

def build_analysis(df: pd.DataFrame, demo_cols: Dict, cfg: Dict, entities: List[str], aliases: Dict, normalizations: List[Dict], expected_raw_cols: int):
    tom_cols = find_cols(df, cfg["tom"])
    esp_cols = find_cols(df, cfg["esp"])
    ayud_cols = find_cols(df, cfg["ayu"])

    raw_cols = [demo_cols.get(k) for k in ["sexo", "edad", "departamento", "estrato", "ingreso"]] + tom_cols + esp_cols + ayud_cols
    raw_cols = list(dict.fromkeys(col for col in raw_cols if col is not None))

    if expected_raw_cols and expected_raw_cols > 0 and len(raw_cols) != expected_raw_cols:
        raise ValueError(f"Se esperaban {expected_raw_cols} columnas, pero se detectaron {len(raw_cols)}. Pon 0 en 'Columnas crudas esperadas'.")

    raw_df = df[raw_cols].copy()
    output_df = raw_df.copy()
    indicators = {}
    tom_response = first_non_empty_response(raw_df, tom_cols)

    aided_map = {entity: [] for entity in entities}
    for col in ayud_cols:
        entity = entity_from_aided_col(col, entities)
        if entity in aided_map: aided_map[entity].append(col)

    for entity in entities:
        tom = tom_response.apply(lambda value: contains_entity(value, entity, aliases)) if tom_cols else pd.Series(False, index=df.index)
        
        espontaneo = pd.Series(False, index=df.index)
        for col in esp_cols:
            espontaneo = espontaneo | raw_df[col].apply(lambda value: contains_entity(value, entity, aliases))

        aided_entity_cols = aided_map.get(entity) or []
        if aided_entity_cols:
            def check_aided_value(value, ent, als):
                if pd.isna(value): return False
                v_str = str(value).strip().lower()
                if v_str in ["1", "si", "sí", "x", "seleccionado", "true"]: return True
                return contains_entity(value, ent, als)
            ayudado = pd.Series(False, index=df.index)
            for col in aided_entity_cols:
                ayudado = ayudado | raw_df[col].apply(lambda value: check_aided_value(value, entity, aliases))
        else:
            ayudado = pd.Series(False, index=df.index)
            for col in ayud_cols:
                ayudado = ayudado | raw_df[col].apply(lambda value: contains_entity(value, entity, aliases))

        indicators[(entity, "TOM")] = tom.astype(int)
        indicators[(entity, "Espontaneo")] = ((~tom) & espontaneo).astype(int)
        indicators[(entity, "Ayudado")] = ((~tom) & (~espontaneo) & ayudado).astype(int)

        output_df[(entity, "TOM")] = indicators[(entity, "TOM")]
        output_df[(entity, "Espontaneo")] = indicators[(entity, "Espontaneo")]
        output_df[(entity, "Ayudado")] = indicators[(entity, "Ayudado")]

    return output_df, pd.DataFrame(indicators), raw_cols


def tom_frequency_table(df: pd.DataFrame, tom_col, normalizations: List[Dict]) -> pd.DataFrame:
    if not tom_col:
        return pd.DataFrame()
    tom_cols = [tom_col] if isinstance(tom_col, str) else list(tom_col)
    tom_response = first_non_empty_response(df, tom_cols)
    
    all_mentions = []
    
    for row_value in tom_response.dropna():
        val_str = str(row_value).strip()
        if val_str.lower() in NEGATIVOS or val_str == "":
            all_mentions.append("Ninguno / NS-NR")
            continue
        
        pieces = re.split(r'[,;]|\s+y\s+|\s+Y\s+', val_str)
        
        for piece in pieces:
            piece = piece.strip()
            if not piece or piece.lower() in NEGATIVOS:
                continue
            
            normalized_piece = apply_normalizations(piece, normalizations)
            v = str(normalized_piece).strip()
            if not (v.startswith("Conglomerado") or v.startswith("Grupo") or v.startswith("Fundación") or v.startswith("Banco")):
                if not v.isupper():
                    v = v.title()
                
            all_mentions.append(v)
            
    if not all_mentions:
        return pd.DataFrame(columns=["Marca / Categoría TOM", "Cantidad", "%"])
        
    freq = pd.Series(all_mentions).value_counts().reset_index()
    freq.columns = ["Marca / Categoría TOM", "Cantidad"]
    total = freq["Cantidad"].sum()
    freq["%"] = freq["Cantidad"] / total if total > 0 else 0
    return freq


def category_table(indicators, df, category_col, category_value, entities):
    mask = pd.Series(True, index=df.index) if category_col is None else df[category_col].fillna("Sin dato") == category_value
    base = int(mask.sum())
    rows = []
    for entity in entities:
        tom = int(indicators[(entity, "TOM")][mask].sum()) if (entity, "TOM") in indicators else 0
        esp = int(indicators[(entity, "Espontaneo")][mask].sum()) if (entity, "Espontaneo") in indicators else 0
        ayu = int(indicators[(entity, "Ayudado")][mask].sum()) if (entity, "Ayudado") in indicators else 0
        awa = tom + esp + ayu
        rows.append({"Entidad": entity, "TOM": tom, "TOM %": tom/base if base else 0, "Espontaneo": esp, "Espontaneo %": esp/base if base else 0, "Ayudado": ayu, "Ayudado %": ayu/base if base else 0, "AWA": awa, "AWA %": awa/base if base else 0})
    return pd.DataFrame(rows)

def awareness_sections(indicators, df, demo_cols, analysis_name, entities):
    sections = [(f"Resumen general {analysis_name}", category_table(indicators, df, None, None, entities))]
    for key in ["sexo", "estrato"]:
        col = demo_cols.get(key)
        if col:
            for value in pd.Series(df[col].fillna("Sin dato")).drop_duplicates():
                sections.append((f"Detalle por {key} - {value}", category_table(indicators, df, col, value, entities)))
    return sections

def department_sections(indicators, df, demo_cols, entities):
    dept_col = demo_cols.get("departamento")
    if not dept_col: return [("Sin departamento/ciudad detectado", pd.DataFrame())]
    return [(str(value), category_table(indicators, df, dept_col, value, entities)) for value in pd.Series(df[dept_col].fillna("Sin dato")).drop_duplicates()]

def demographic_sections(df, demo_cols):
    sections = [("Base", pd.DataFrame({"Indicador": ["Total encuestas procesadas"], "Valor": [len(df)]}))]
    for key, label in [("sexo", "Sexo"), ("edad", "Edad"), ("departamento", "Ciudad/Depto"), ("estrato", "Estrato"), ("ingreso", "Ingreso")]:
        col = demo_cols.get(key)
        if col:
            table = df[col].fillna("Sin dato").value_counts().rename_axis(label).reset_index(name="Cantidad")
            table["%"] = table["Cantidad"] / len(df) if len(df) else 0
            sections.append((f"Demográfico - {label}", table))
    return sections

def style_sheet(ws):
    border = Border(left=Side(style="thin", color="BFBFBF"), right=Side(style="thin", color="BFBFBF"), top=Side(style="thin", color="BFBFBF"), bottom=Side(style="thin", color="BFBFBF"))
    for row in ws.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col_idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24

def write_main_sheet(wb, sheet_name, output_df):
    ws = wb.create_sheet(safe_sheet_name(sheet_name))
    for col_idx, column in enumerate(output_df.columns, start=1):
        ws.cell(1, col_idx).value = str(column)
        ws.cell(1, col_idx).fill = PatternFill("solid", fgColor="D9EAF7")
        ws.cell(1, col_idx).font = Font(bold=True)
    for row_idx, row in enumerate(output_df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row_idx, col_idx).value = value
    style_sheet(ws)

def write_sections_sheet(wb, sheet_name, sections):
    ws = wb.create_sheet(safe_sheet_name(sheet_name))
    row_idx = 1
    for title, df_table in sections:
        ws.cell(row_idx, 1).value = title
        ws.cell(row_idx, 1).fill = PatternFill("solid", fgColor="C6E0B4")
        ws.cell(row_idx, 1).font = Font(bold=True)
        row_idx += 1
        if df_table.empty:
            ws.cell(row_idx, 1).value = "Sin datos disponibles"
            row_idx += 3
            continue
        for col_idx, column in enumerate(df_table.columns, start=1):
            ws.cell(row_idx, col_idx).value = column
            ws.cell(row_idx, col_idx).fill = PatternFill("solid", fgColor="4472C4")
            ws.cell(row_idx, col_idx).font = Font(bold=True, color="FFFFFF")
        row_idx += 1
        for record in df_table.itertuples(index=False):
            for col_idx, value in enumerate(record, start=1):
                cell = ws.cell(row_idx, col_idx)
                cell.value = value
                if str(df_table.columns[col_idx - 1]).endswith("%"):
                    cell.number_format = "0%"
            row_idx += 1
        row_idx += 2
    style_sheet(ws)

def build_excel_bytes(df, demo_cols, cfg1, cfg2, entities, aliases, normalizations, expected_raw_cols):
    run_1 = any(cfg1[k].strip() not in ["0", ""] for k in ["tom", "esp", "ayu"])
    run_2 = any(cfg2[k].strip() not in ["0", ""] for k in ["tom", "esp", "ayu"])
    
    if not run_1 and not run_2: 
        raise ValueError("Debe configurar al menos un análisis válido (TOM, Espontáneo o Ayudado). Escriba los prefijos de columnas en la sección 2.")
    
    wb = Workbook()
    wb.remove(wb.active)
    write_sections_sheet(wb, "Resumen Demográficos", demographic_sections(df, demo_cols))

    for cfg, run_flag in [(cfg1, run_1), (cfg2, run_2)]:
        if run_flag:
            output_df, indicators, _ = build_analysis(df, demo_cols, cfg, entities, aliases, normalizations, expected_raw_cols)
            write_main_sheet(wb, cfg["name"], output_df)
            write_sections_sheet(wb, f"Resumen {cfg['name']}", awareness_sections(indicators, df, demo_cols, cfg['name'], entities))
            write_sections_sheet(wb, f"Deptos {cfg['name']}", department_sections(indicators, df, demo_cols, entities))
            
            tom_col_names = find_cols(df, cfg["tom"])
            if tom_col_names:
                freq_table = tom_frequency_table(df, tom_col_names, normalizations)
                write_sections_sheet(wb, f"Frecuencias P1 {cfg['name']}", [("Distribución de Respuestas P1 (TOM)", freq_table)])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
