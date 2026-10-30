# -*- coding: utf-8 -*-
"""
core/export.py
Reporte en Excel: datos codificados, frecuencias e insights (reutiliza el estilo de awareness).
"""

import io
import re
from typing import Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from core.awareness import safe_sheet_name, write_main_sheet, write_sections_sheet
from core.coding import NSNR_LABEL, frequency_table, onehot, valid_base

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def short_name(text: str, size: int = 40) -> str:
    text = re.sub(r"\s+", " ", str(text)).strip()
    return text if len(text) <= size else text[: size - 1] + "…"


def _clean(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if not pd.api.types.is_numeric_dtype(out[col]):
            out[col] = out[col].map(lambda v: _ILLEGAL.sub("", v) if isinstance(v, str) else v)
    return out


def build_report_excel(df: pd.DataFrame, coded: Dict[str, Dict], insights: Dict[str, Dict],
                       validation: Optional[Dict[str, Dict]] = None) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # 1) Datos originales + columnas codificadas
    out = df.copy().reset_index(drop=True)
    for n, (q, info) in enumerate(coded.items(), start=1):
        tag = f"P{n}"
        labels = info["labels"]
        out[f"{tag} | Categorías"] = [" | ".join(l) for l in labels]
        oh = onehot(labels, info["categories"])
        oh.columns = [f"{tag} · {short_name(c, 25)}" for c in oh.columns]
        out = pd.concat([out, oh], axis=1)
    write_main_sheet(wb, "Datos codificados", _clean(out))

    # 2) Frecuencias
    sections = []
    for n, (q, info) in enumerate(coded.items(), start=1):
        table = frequency_table(info["labels"], info["categories"])
        base = valid_base(info["labels"])
        sections.append((f"P{n} · {short_name(q, 80)}  (base: {base} con respuesta válida)", _clean(table)))
    if sections:
        write_sections_sheet(wb, "Frecuencias", sections)

    # 3) Insights
    ins_sections = []
    for n, (q, info) in enumerate(coded.items(), start=1):
        ins = insights.get(q)
        if not ins:
            continue
        rows = [{
            "Categoría": c["categoria"], "Personas": c["n"], "%": c["pct"],
            "Insight": c["explicacion"], "Cita 1": (c["citas"] + ["", ""])[0], "Cita 2": (c["citas"] + ["", ""])[1],
        } for c in ins["categorias"]]
        ins_sections.append((f"P{n} · {short_name(q, 80)}", pd.DataFrame(
            [{"Categoría": "Resumen", "Personas": "", "%": "", "Insight": ins["resumen"], "Cita 1": "", "Cita 2": ""}] + rows)))
    if ins_sections:
        write_sections_sheet(wb, "Insights", [(t, _clean(d)) for t, d in ins_sections])
        ws = wb["Insights"]
        for letter, width in (("A", 28), ("B", 10), ("C", 8), ("D", 70), ("E", 45), ("F", 45)):
            ws.column_dimensions[letter].width = width

    # 4) Validación de segunda opinión (si se ejecutó)
    val_sections = []
    for n, q in enumerate(coded, start=1):
        result = (validation or {}).get(q, {}).get("second")
        if result and not result["metrics"].empty:
            val_sections.append((f"P{n} · {short_name(q, 80)} · Métricas por categoría", _clean(result["metrics"])))
            if not result["review"].empty:
                val_sections.append((f"P{n} · Casos para revisión humana", _clean(result["review"].head(1000))))
    if val_sections:
        write_sections_sheet(wb, "Validación", val_sections)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_study_categorization_excel(df: pd.DataFrame, classifications: list[Dict]) -> bytes:
    """Exporta tres hojas limpias y evita asignar etiquetas a respuestas vacías."""
    from openpyxl.utils import get_column_letter
    from core.coding import is_blank_response

    wb = Workbook(write_only=True)

    def safe_value(value):
        if value is None:
            return None
        if hasattr(value, "item"):
            try:
                value = value.item()
            except (ValueError, AttributeError):
                pass
        try:
            if pd.isna(value):
                return None
        except (TypeError, ValueError):
            pass
        if isinstance(value, str):
            if not value.strip():
                return None
            return _ILLEGAL.sub("", value)[:32767]
        return value

    def append_header(ws, headers, widths=None):
        cells = []
        for value in headers:
            cell = WriteOnlyCell(ws, value=safe_value(str(value)))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cells.append(cell)
        ws.append(cells)
        ws.freeze_panes = "A2"
        if widths:
            for index, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(index)].width = width
        return len(headers)

    # Hoja 1: copia legible de los datos importados, sin columnas de resultados.
    data_ws = wb.create_sheet(safe_sheet_name("Datos"))
    data_columns = [str(column) for column in df.columns]
    data_count = append_header(data_ws, data_columns, [22] * len(data_columns))
    for row in df.itertuples(index=False, name=None):
        data_ws.append([safe_value(value) for value in row])
    if data_count:
        data_ws.auto_filter.ref = f"A1:{get_column_letter(data_count)}{len(df) + 1}"

    # Hoja 2: una pareja respuesta/categoría por pregunta y por persona.
    result_ws = wb.create_sheet(safe_sheet_name("Preguntas y categorización"))
    result_headers = ["Fila de encuesta"]
    widths = [16]
    for item in classifications:
        question = str(item.get("question", "Pregunta"))
        result_headers.extend([f"{question} | Respuesta", f"{question} | Categoría(s)"])
        widths.extend([42, 30])
    result_count = append_header(result_ws, result_headers, widths)
    columns = list(df.columns)
    responses_by_question = {
        item.get("question", ""): df[item.get("question")].tolist() if item.get("question") in columns else [None] * len(df)
        for item in classifications
    }
    for row_index in range(len(df)):
        values = [row_index + 2]  # fila equivalente en la hoja Datos (encabezado en fila 1)
        for item in classifications:
            question = item.get("question", "")
            response = responses_by_question[question][row_index]
            saved_labels = item.get("labels") or []
            labels = saved_labels[row_index] if row_index < len(saved_labels) else []
            if is_blank_response(response):
                values.extend([None, None])
            else:
                values.extend([safe_value(response), safe_value(" | ".join(str(label) for label in (labels or [])))])
        result_ws.append(values)
    if result_count:
        result_ws.auto_filter.ref = f"A1:{get_column_letter(result_count)}{len(df) + 1}"

    # Hoja 3: codebook, sin repetir categorías, términos ni alias.
    dictionary_ws = wb.create_sheet(safe_sheet_name("Diccionario"))
    dictionary_headers = ["Pregunta", "Tipo de análisis", "Código / categoría", "Descripción", "Palabras clave / sinónimos", "Alias / variantes"]
    dictionary_count = append_header(dictionary_ws, dictionary_headers, [24, 20, 30, 62, 42, 42])
    kind_names = {"brands": "Marcas", "sentiment": "Sentimiento", "themes": "Temas"}
    dictionary_rows = 1
    for item in classifications:
        question = str(item.get("question", ""))
        kind = str(item.get("kind", ""))
        kind_label = kind_names.get(kind, kind)
        frame = [entry for entry in (item.get("codeframe") or []) if isinstance(entry, dict)]
        frame_by_name = {str(entry.get("nombre", "")): entry for entry in frame}
        config = item.get("config") or {}
        categories = list(item.get("categories") or [])
        if kind == "brands":
            categories = list(dict.fromkeys([*(config.get("known_brands") or []), *categories, "Otra", "Ninguna", "NS/NR"]))
        else:
            categories = list(dict.fromkeys([*categories, "Otros", "NS/NR"]))
        brand_mapping = item.get("brand_mapping") or {}
        for category in categories:
            name = str(category)
            entry = frame_by_name.get(name, {})
            keywords = entry.get("palabras_clave", entry.get("sinonimos", []))
            if isinstance(keywords, list):
                keywords = ", ".join(str(value) for value in keywords)
            aliases = []
            if kind == "brands":
                aliases = [str(raw) for raw, canonical in brand_mapping.items() if str(canonical) == name and str(raw) != name]
                if not aliases:
                    aliases = (item.get("aliases") or {}).get(name, [])
            description = entry.get("descripcion", "")
            if not description:
                description = {
                    "Otra": "Mención de marca fuera del catálogo permitido.",
                    "Ninguna": "Respuesta explícita que indica que no menciona una marca.",
                    "Otros": "Respuesta válida que no encaja en las categorías definidas.",
                    "NS/NR": "Respuesta explícita sin información; una celda vacía permanece en blanco.",
                }.get(name, "Marca del catálogo permitido." if kind == "brands" else "")
            dictionary_ws.append([
                safe_value(question), kind_label, safe_value(name), safe_value(description),
                safe_value(str(keywords or "")), safe_value(" | ".join(str(value) for value in aliases)),
            ])
            dictionary_rows += 1
    if dictionary_count:
        dictionary_ws.auto_filter.ref = f"A1:{get_column_letter(dictionary_count)}{dictionary_rows}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
